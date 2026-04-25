from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def clean_profile_text(value: Any) -> str | None:
    text = clean_text(value)
    if not text:
        return None
    # Drop numeric placeholders from enrichment sheets (e.g. "17").
    if text.isdigit():
        return None
    return text


def parse_year(value: Any) -> int | None:
    if value is None:
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        text = clean_text(value)
        if not text:
            return None
        digits = "".join(ch for ch in text if ch.isdigit())
        if len(digits) >= 4:
            return int(digits[:4])
        return None


def build_record(headers: list[str], row: tuple[Any, ...], updated_on: str) -> dict[str, Any] | None:
    values = {headers[idx]: row[idx] if idx < len(row) else None for idx in range(len(headers))}

    name = clean_text(values.get("Nama"))
    if not name:
        return None

    graduation_year = parse_year(values.get("Angkatan"))
    current_title = clean_profile_text(values.get("Current_Title"))
    current_company = clean_profile_text(values.get("Current_Company_or_Institution"))
    industry = clean_profile_text(values.get("Industry"))
    location = clean_profile_text(values.get("Location"))
    highest_education = clean_profile_text(values.get("Highest_Education"))
    confidence = clean_text(values.get("Confidence"))
    notes = clean_text(values.get("Notes"))
    profile_url = clean_text(values.get("Public_Source_URL"))

    record: dict[str, Any] = {
        "name": name,
        "headline": current_title,
        "location": location,
        "connections": "",
        "about": notes,
        "experiences": [],
        "education": [],
        "skills": [],
        "certifications": [],
        "languages": [],
        "projects": [],
        "volunteer_experience": [],
        "contact_info": {
            "email": "",
            "phone": "",
            "website": "",
            "linkedin_url": profile_url or "",
        },
        "last_updated": updated_on,
        "graduation_year": graduation_year,
        "industry": industry,
        "current_company": current_company,
        "confidence": confidence,
        "source_batches": clean_text(values.get("Source_Batches")),
        "source_url": profile_url,
        "merge_key": clean_text(values.get("Merge_Key")),
    }

    if current_title or current_company:
        record["experiences"].append(
            {
                "position": current_title,
                "company": current_company,
                "duration": "",
                "location": location,
                "description": "",
            }
        )

    if highest_education:
        record["education"].append(
            {
                "school": highest_education,
                "degree": "",
                "field_of_study": "",
                "duration": str(graduation_year) if graduation_year else "",
            }
        )

    return record


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert alumni master Excel into data-alumni.json format.")
    parser.add_argument("--input", required=True, help="Path to alumni_master_merged.xlsx")
    parser.add_argument("--output", required=True, help="Path to output JSON file")
    parser.add_argument("--sheet", default="Master_Unique", help="Worksheet name containing normalized alumni rows")
    args = parser.parse_args()

    workbook = load_workbook(args.input, data_only=True)
    if args.sheet not in workbook.sheetnames:
        raise ValueError(f"Sheet '{args.sheet}' not found. Available: {workbook.sheetnames}")

    worksheet = workbook[args.sheet]
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [clean_text(cell) or f"column_{idx}" for idx, cell in enumerate(header_row)]

    updated_on = date.today().isoformat()
    records: list[dict[str, Any]] = []
    seen_keys: set[str] = set()

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        record = build_record(headers, row, updated_on)
        if not record:
            continue
        merge_key = record.get("merge_key") or record["name"].upper()
        if merge_key in seen_keys:
            continue
        seen_keys.add(merge_key)
        records.append(record)

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {len(records)} records to {output_path}")


if __name__ == "__main__":
    main()